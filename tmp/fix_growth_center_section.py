from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph


ROOT = Path(r"D:\Project Files\GEOFlow")
DOCX = ROOT / "docs" / "GEOFlow普通管理员使用说明-图文版.docx"
OUTPUT = DOCX
SCREENSHOT_DIR = ROOT / "docs" / "manual-screenshots"


def insert_after(paragraph: Paragraph, text: str = "", style: str | None = None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    inserted = Paragraph(new_p, paragraph._parent)
    if style:
        inserted.style = style
    if text:
        inserted.add_run(text)
    return inserted


def delete_paragraph(paragraph: Paragraph) -> None:
    p = paragraph._element
    p.getparent().remove(p)
    paragraph._p = paragraph._element = None


def bold_runs(paragraph: Paragraph) -> None:
    for run in paragraph.runs:
        run.bold = True


def caption_style(paragraph: Paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(10)
    for run in paragraph.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(71, 85, 105)


def image_style(paragraph: Paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(2)


def add_text(current: Paragraph, text: str, style: str | None = None, bold: bool = False) -> Paragraph:
    current = insert_after(current, text, style)
    if bold:
        bold_runs(current)
    return current


def add_image(current: Paragraph, file_name: str, caption: str) -> Paragraph:
    image_p = insert_after(current)
    image_style(image_p)
    image_p.add_run().add_picture(str(SCREENSHOT_DIR / file_name), width=Inches(6.25))
    caption_p = insert_after(image_p, caption)
    caption_style(caption_p)
    return caption_p


def main() -> None:
    doc = Document(str(DOCX))
    paragraphs = list(doc.paragraphs)

    start = end = None
    for i, paragraph in enumerate(paragraphs):
        text = paragraph.text.strip()
        if text == "八、增长中心":
            start = i
        elif start is not None and text == "九、任务管理":
            end = i
            break

    if start is None or end is None:
        raise RuntimeError("Cannot locate growth center section.")

    anchor = paragraphs[start]
    for paragraph in paragraphs[start + 1 : end]:
        delete_paragraph(paragraph)

    current = anchor
    current = add_text(current, "模块作用：增长中心不只是查看数据的地方，也负责线索收集和线索跟进。你可以在这里查看访问、AI 爬虫迹象、内容表现、分发效果，也可以创建转化表单、管理表单，并处理客户提交的线索。")
    current = add_text(current, "前后关系：增长中心通常在文章发布后使用。文章产生访问后，你可以通过转化表单收集咨询、报名、试用申请等线索，再在线索收件箱中跟进处理。")
    current = add_text(current, "使用结果：你能看到内容运营数据，也能把前台访问转成可跟进的客户线索。")
    current = add_image(current, "08-增长中心-总览.png", "图：增长中心总览，可进入创建表单、表单管理、线索收件箱，并查看整体增长数据。")

    current = add_text(current, "查看全局数据", "Heading 2")
    current = add_text(current, "功能说明：通过总览判断访问、线索、内容生产、任务执行和分发是否正常。")
    current = add_text(current, "操作步骤：", bold=True)
    current = add_text(current, "1. 点击左侧导航“增长中心”。")
    current = add_text(current, "2. 先查看页面顶部的访问、AI 访问、线索提交、新线索、待跟进等数据。")
    current = add_text(current, "3. 再查看内容、任务、分发、访问日志等运营数据。")
    current = add_text(current, "4. 如果看到新线索或待跟进数量，优先进入线索收件箱处理。")
    current = add_text(current, "完成后检查：", bold=True)
    current = add_text(current, "• 能判断当前是否有新线索、待跟进线索或异常运营数据。")
    current = add_text(current, "注意事项：", bold=True)
    current = add_text(current, "• 增长中心展示的是运营数据，不等同于搜索排名承诺。")

    current = add_text(current, "创建线索表单", "Heading 2")
    current = add_text(current, "功能说明：创建前台可提交的转化表单，用于收集客户咨询、联系方式、需求说明等信息。系统页面中称为“转化表单”。")
    current = add_text(current, "使用前准备：", bold=True)
    current = add_text(current, "• 已明确要收集哪些信息，例如姓名、手机号、邮箱、需求说明。")
    current = add_image(current, "08-增长中心-创建线索表单.png", "图：创建转化表单页面，可设置表单名称、访问标识、状态、按钮文案和字段。")
    current = add_text(current, "操作步骤：", bold=True)
    current = add_text(current, "1. 在增长中心点击“创建表单”。")
    current = add_text(current, "2. 填写表单名称，名称要方便后台识别。")
    current = add_text(current, "3. 设置 Slug。Slug 会生成前台访问地址，建议使用英文、数字或短横线。")
    current = add_text(current, "4. 设置状态。需要前台可用时选择“启用”。")
    current = add_text(current, "5. 设置按钮文案，例如“提交”“立即咨询”。")
    current = add_text(current, "6. 在字段配置中添加需要客户填写的字段。")
    current = add_text(current, "7. 检查无误后保存。")
    current = add_text(current, "完成后检查：", bold=True)
    current = add_text(current, "• 表单保存后能在表单管理列表中看到。")
    current = add_text(current, "注意事项：", bold=True)
    current = add_text(current, "• 不要收集不必要的敏感信息，字段越少，客户越容易提交。")

    current = add_text(current, "管理线索表单", "Heading 2")
    current = add_text(current, "功能说明：查看已有表单、确认表单是否启用，并对表单进行编辑或停用。")
    current = add_image(current, "08-增长中心-表单管理.png", "图：表单管理页面，用于查看、编辑和管理已创建的转化表单。")
    current = add_text(current, "操作步骤：", bold=True)
    current = add_text(current, "1. 在增长中心点击“表单管理”。")
    current = add_text(current, "2. 查看表单名称、访问标识、状态和提交数量。")
    current = add_text(current, "3. 需要调整字段或文案时，进入编辑。")
    current = add_text(current, "4. 暂时不想让客户提交时，将表单停用。")
    current = add_text(current, "完成后检查：", bold=True)
    current = add_text(current, "• 需要使用的表单处于启用状态，不需要使用的表单已停用。")
    current = add_text(current, "注意事项：", bold=True)
    current = add_text(current, "• 表单已对外使用后，修改字段前要确认不会影响正在投放的页面。")

    current = add_text(current, "查看和处理线索", "Heading 2")
    current = add_text(current, "功能说明：线索收件箱用于查看客户提交的信息，并记录跟进状态。")
    current = add_image(current, "08-增长中心-线索收件箱.png", "图：线索收件箱页面，用于查看客户提交记录、筛选状态并导出线索。")
    current = add_text(current, "操作步骤：", bold=True)
    current = add_text(current, "1. 在增长中心点击“线索收件箱”。")
    current = add_text(current, "2. 查看新线索、已联系、有效、已转化等状态。")
    current = add_text(current, "3. 点击线索详情，查看客户提交的具体内容。")
    current = add_text(current, "4. 跟进客户后，按实际情况更新线索状态。")
    current = add_text(current, "5. 需要交给销售或运营同事时，可以导出线索。")
    current = add_text(current, "完成后检查：", bold=True)
    current = add_text(current, "• 新线索被及时处理，待跟进数量减少。")
    current = add_text(current, "注意事项：", bold=True)
    current = add_text(current, "• 线索状态要及时更新，否则后续复盘时会分不清哪些已经联系、哪些还没处理。")

    current = add_text(current, "分析内容表现", "Heading 2")
    current = add_text(current, "功能说明：找出访问较高、主题较好的文章，为后续选题提供依据。")
    current = add_text(current, "使用前准备：", bold=True)
    current = add_text(current, "• 至少已经发布过文章，并产生访问记录。")
    current = add_text(current, "操作步骤：", bold=True)
    current = add_text(current, "1. 查看热门文章访问区域。")
    current = add_text(current, "2. 记录访问量较高的标题、分类和主题。")
    current = add_text(current, "3. 进入内容管理查看这些文章的结构和关键词。")
    current = add_text(current, "4. 把表现好的主题补充到标题库或选题规划中。")
    current = add_text(current, "完成后检查：", bold=True)
    current = add_text(current, "• 能得到下一批内容选题方向。")
    current = add_text(current, "注意事项：", bold=True)
    current = add_text(current, "• 访问少不一定代表文章无效，新站和新文章需要时间积累。")

    current = add_text(current, "观察分发效果", "Heading 2")
    current = add_text(current, "功能说明：判断文章是否成功同步到目标站点，及时处理失败队列。")
    current = add_text(current, "使用前准备：", bold=True)
    current = add_text(current, "• 需要已经配置分发渠道，并有文章触发分发。")
    current = add_text(current, "操作步骤：", bold=True)
    current = add_text(current, "1. 在增长中心查看分发状态概览。")
    current = add_text(current, "2. 重点看失败、待分发、已同步数量。")
    current = add_text(current, "3. 失败数量不为 0 时，进入分发管理的分发队列。")
    current = add_text(current, "4. 查看错误原因后修复渠道或重试。")
    current = add_text(current, "完成后检查：", bold=True)
    current = add_text(current, "• 失败队列减少，已同步数量增加。")
    current = add_text(current, "注意事项：", bold=True)
    current = add_text(current, "• 同一渠道持续失败时，不要反复重试，先检查渠道配置。")

    doc.save(str(OUTPUT))
    print(str(OUTPUT))


if __name__ == "__main__":
    main()
