from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph


ROOT = Path(r"D:\Project Files\GEOFlow")
SOURCE = ROOT / "docs" / "GEOFlow普通管理员使用说明.docx"
OUTPUT = ROOT / "docs" / "GEOFlow普通管理员使用说明-图文版.docx"
SCREENSHOT_DIR = ROOT / "docs" / "manual-screenshots"


INSERTIONS = {
    "七、首页": [
        ("07-首页-总览.png", "图：后台首页总览，进入系统后先看快捷入口和整体运行状态。"),
    ],
    "八、增长中心": [
        ("08-增长中心-总览.png", "图：增长中心总览，用于查看访问、线索和内容效果数据。"),
    ],
    "九、任务管理": [
        ("09-任务管理-列表.png", "图：任务管理列表，可查看任务状态，也可以进入新建任务页面。"),
    ],
    "创建任务基础信息": [
        ("09-任务管理-新建基础信息.png", "图：新建任务页面上方区域，先确认任务入口、任务包说明、任务名称和标题库选择。"),
    ],
    "十、分发管理": [
        ("10-分发管理-渠道列表.png", "图：分发管理列表，用于查看和维护外部发布目标。"),
    ],
    "创建分发渠道": [
        ("10-分发管理-创建渠道.png", "图：创建分发渠道页面，按目标站点类型填写连接信息。"),
    ],
    "查看分发队列": [
        ("10-分发管理-分发队列.png", "图：分发队列页面，用于查看文章同步到外部站点的执行结果。"),
    ],
    "十一、内容管理": [
        ("11-内容管理-文章列表.png", "图：内容管理列表，可筛选、审核、编辑和发布文章。"),
    ],
    "编辑文章": [
        ("11-内容管理-文章编辑.png", "图：文章编辑页面，用于维护标题、正文、分类、状态等内容。"),
    ],
    "十二、知识资产": [
        ("12-知识资产-总览.png", "图：知识资产总览，统一进入知识库、标题库、关键词库、图片库和作者管理。"),
    ],
    "创建知识库": [
        ("12-知识资产-创建知识库.png", "图：创建知识库页面，先建立知识库，再上传企业资料。"),
    ],
    "维护关键词库": [
        ("12-知识资产-关键词库.png", "图：关键词库列表，用于维护任务和选题会用到的关键词素材。"),
    ],
    "维护标题库": [
        ("12-知识资产-标题库.png", "图：标题库列表，用于维护文章生成时可以调用的标题素材。"),
    ],
    "维护图片库和作者": [
        ("12-知识资产-图片库.png", "图：图片库列表，用于维护文章配图素材。"),
    ],
    "十三、关键词趋势": [
        ("13-关键词趋势-列表.png", "图：关键词趋势列表，用于管理趋势来源和抓取结果。"),
    ],
    "创建趋势数据源": [
        ("13-关键词趋势-创建来源.png", "图：创建趋势数据源页面，用于设置关键词来源和抓取规则。"),
    ],
    "十四、谷歌搜录": [
        ("14-谷歌搜录-总览.png", "图：谷歌搜录总览，用于查看 Google Search Console 连接和站点数据。"),
    ],
    "十五、选题规划": [
        ("15-选题规划-列表.png", "图：选题规划列表，用于查看已生成的选题计划。"),
    ],
    "创建选题计划": [
        ("15-选题规划-创建计划.png", "图：创建选题计划页面，选择关键词、知识库和生成要求后生成候选选题。"),
    ],
    "十六、AI 配置器": [
        ("16-AI配置器-总览.png", "图：AI 配置器总览，用于进入模型、提示词和特殊提示词配置。"),
    ],
    "添加并测试 AI 模型": [
        ("16-AI配置器-模型配置.png", "图：AI 模型配置页面，用于查看已有模型、添加模型和测试连接。"),
    ],
    "设置提示词": [
        ("16-AI配置器-提示词.png", "图：提示词配置页面，用于维护文章生成时使用的写作要求。"),
    ],
    "设置特殊提示词": [
        ("16-AI配置器-特殊提示词.png", "图：特殊提示词页面，用于维护关键词、描述等专项生成规则。"),
    ],
}


def insert_paragraph_after(paragraph: Paragraph) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    return Paragraph(new_p, paragraph._parent)


def format_caption(paragraph: Paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(10)
    for run in paragraph.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(71, 85, 105)


def format_image_paragraph(paragraph: Paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(2)


def find_anchor(doc: Document, heading_text: str) -> Paragraph:
    paragraphs = list(doc.paragraphs)
    for idx, paragraph in enumerate(paragraphs):
        if paragraph.text.strip() != heading_text:
            continue

        style_name = paragraph.style.name if paragraph.style else ""
        if style_name == "Heading 1":
            anchor = paragraph
            for next_paragraph in paragraphs[idx + 1 :]:
                next_style = next_paragraph.style.name if next_paragraph.style else ""
                if next_style.startswith("Heading"):
                    break
                if next_paragraph.text.strip():
                    anchor = next_paragraph
            return anchor

        return paragraph

    raise ValueError(f"Cannot find heading: {heading_text}")


def main() -> None:
    doc = Document(str(SOURCE))

    for heading, images in INSERTIONS.items():
        anchor = find_anchor(doc, heading)
        current = anchor
        for file_name, caption in images:
            image_path = SCREENSHOT_DIR / file_name
            if not image_path.exists():
                raise FileNotFoundError(image_path)

            image_paragraph = insert_paragraph_after(current)
            format_image_paragraph(image_paragraph)
            image_paragraph.add_run().add_picture(str(image_path), width=Inches(6.25))

            caption_paragraph = insert_paragraph_after(image_paragraph)
            caption_paragraph.add_run(caption)
            format_caption(caption_paragraph)

            current = caption_paragraph

    doc.save(str(OUTPUT))
    print(str(OUTPUT))


if __name__ == "__main__":
    main()
