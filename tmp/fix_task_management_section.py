from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph


ROOT = Path(r"D:\Project Files\GEOFlow")
DOCX = ROOT / "docs" / "GEOFlow普通管理员使用说明-图文版.docx"
OUTPUT = DOCX
TASK_CREATE_IMAGE = ROOT / "docs" / "manual-screenshots" / "09-任务管理-新建基础信息.png"


def insert_after(paragraph: Paragraph, text: str = "", style: str | None = None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    inserted = Paragraph(new_p, paragraph._parent)
    if style:
        inserted.style = style
    if text:
        inserted.add_run(text)
    return inserted


def delete_paragraph(paragraph: Paragraph) -> None:
    p = paragraph._element
    p.getparent().remove(p)
    paragraph._p = paragraph._element = None


def format_caption(paragraph: Paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(10)
    for run in paragraph.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(71, 85, 105)


def format_image_paragraph(paragraph: Paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(2)


def format_label(paragraph: Paragraph) -> None:
    for run in paragraph.runs:
        run.bold = True


def add_lines(anchor: Paragraph, lines: list[tuple[str, str | None, str | None]]) -> Paragraph:
    current = anchor
    for text, style, role in lines:
        current = insert_after(current, text, style)
        if role == "label":
            format_label(current)
        elif role == "caption":
            format_caption(current)
        elif role == "spacer":
            current.paragraph_format.space_after = Pt(4)
    return current


def main() -> None:
    doc = Document(str(DOCX))
    paragraphs = list(doc.paragraphs)

    start = end = None
    for i, paragraph in enumerate(paragraphs):
        text = paragraph.text.strip()
        if text == "九、任务管理":
            start = i
        elif start is not None and text == "十、分发管理":
            end = i
            break

    if start is None or end is None:
        raise RuntimeError("Cannot locate task management section.")

    # Keep Heading 1 + three module intro paragraphs + the existing task-list overview image.
    anchor = paragraphs[start + 4]

    for paragraph in paragraphs[start + 5 : end]:
        delete_paragraph(paragraph)

    lines: list[tuple[str, str | None, str | None]] = [
        ("图：任务管理列表，可查看任务状态，也可以进入新建任务页面。", None, "caption"),
        ("查看任务列表", "Heading 2", None),
        ("功能说明：任务列表用于查看已有任务、确认任务是否启用，以及观察任务生成文章的情况。", None, None),
        ("操作步骤：", None, "label"),
        ("1. 点击左侧导航“任务管理”。", None, None),
        ("2. 在任务列表中查看任务名称、状态、已创建文章、草稿、已发布等信息。", None, None),
        ("3. 需要新增任务时，点击页面中的“新建任务”或“创建任务”入口。", None, None),
        ("4. 需要调整已有任务时，进入对应任务的编辑入口。", None, None),
        ("完成后检查：", None, "label"),
        ("• 能判断当前任务是否启用，以及任务是否已经产生文章。", None, None),
        ("注意事项：", None, "label"),
        ("• 任务列表是管理入口，不是文章审核入口；文章内容需要到“内容管理”里处理。", None, None),
        ("创建任务", "Heading 2", None),
        ("功能说明：创建任务是一个完整流程，不是多个互相独立的小功能。你需要在同一个创建页面里依次完成基础检查、基础信息、内容生成、图片审核、发布范围和保存检查。", None, None),
        ("使用前准备：", None, "label"),
        ("• AI 配置器里至少有一个可用的聊天模型。", None, None),
        ("• AI 配置器里已经准备好内容提示词。", None, None),
        ("• 知识资产里已经准备标题库；如果需要引用企业资料，还要准备知识库。", None, None),
        ("• 如果任务要配图，需要先准备图片库。", None, None),
        ("• 如果任务要同步到外部站点，需要先准备分发渠道。", None, None),
    ]

    current = add_lines(anchor, lines)

    image_paragraph = insert_after(current)
    format_image_paragraph(image_paragraph)
    image_paragraph.add_run().add_picture(str(TASK_CREATE_IMAGE), width=Inches(6.25))

    current = insert_after(image_paragraph, "图：创建任务页面上方区域，先确认任务入口、任务包说明、任务名称和标题库选择。")
    format_caption(current)

    more_lines: list[tuple[str, str | None, str | None]] = [
        ("操作步骤：", None, "label"),
        ("1. 点击左侧导航“任务管理”。", None, None),
        ("2. 点击“新建任务”或“创建任务”，进入创建页面。", None, None),
        ("3. 先检查页面上的标题库、提示词、AI 模型等下拉项是否有可选内容。没有可选内容时，先回到对应模块补配置。", None, None),
        ("4. 在“任务名称”中填写清楚的名称，例如“产品 FAQ 文章生成-7月”。", None, None),
        ("5. 在“标题库”中选择本次任务要使用的标题来源。", None, None),
        ("6. 在内容配置区域选择内容提示词和 AI 模型。普通使用选择固定模型即可。", None, None),
        ("7. 如果需要引用企业资料，选择和本次主题相关的知识库。知识库不是越多越好，优先选择最相关的资料。", None, None),
        ("8. 如果需要配图，选择图片库并设置每篇文章使用的图片数量。", None, None),
        ("9. 根据实际情况设置是否需要审核。新客户建议先开启审核，不建议一开始全自动发布。", None, None),
        ("10. 设置发布范围：只发布到本站，或同时进入外部分发流程。", None, None),
        ("11. 如果选择外部分发，勾选要使用的分发渠道。", None, None),
        ("12. 检查必填项无误后点击保存。", None, None),
        ("完成后检查：", None, "label"),
        ("• 回到任务列表后，能看到新创建的任务。", None, None),
        ("• 任务状态、标题库、模型、发布范围等信息符合预期。", None, None),
        ("• 启用状态下，系统会按调度执行；暂停状态下，不会立即自动生成。", None, None),
        ("注意事项：", None, "label"),
        ("• 新手测试建议先创建小任务，确认文章质量和发布流程后，再扩大生成数量。", None, None),
        ("• 下拉框为空时不要继续保存，先回到 AI 配置器或知识资产补齐基础资料。", None, None),
        ("• 全自动发布前必须先确认文章质量，否则容易把未检查内容公开。", None, None),
    ]
    add_lines(current, more_lines)

    doc.save(str(OUTPUT))
    print(str(OUTPUT))


if __name__ == "__main__":
    main()
