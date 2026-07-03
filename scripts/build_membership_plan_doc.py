from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "GEOFlow会员系统方案-仅文章知识库限制.docx"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(10)


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "D9E2EC")


def add_heading(doc: Document, text: str, level: int) -> None:
    paragraph = doc.add_heading(text, level=level)
    for run in paragraph.runs:
        run.font.name = "Microsoft YaHei"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.color.rgb = RGBColor(46, 116, 181)


def add_body(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(text)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.1
    for run in paragraph.runs:
        run.font.name = "Microsoft YaHei"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.size = Pt(11)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(4)
        run = paragraph.add_run(item)
        run.font.name = "Microsoft YaHei"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.size = Pt(10.5)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)

    header_cells = table.rows[0].cells
    for idx, header in enumerate(headers):
        set_cell_shading(header_cells[idx], "F2F4F7")
        set_cell_text(header_cells[idx], header, bold=True)
        header_cells[idx].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        if widths:
            header_cells[idx].width = Inches(widths[idx])

    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], value)
            cells[idx].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if widths:
                cells[idx].width = Inches(widths[idx])

    doc.add_paragraph()


def add_code_block(doc: Document, lines: list[str]) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, "F7F9FC")
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    for index, line in enumerate(lines):
        if index > 0:
            paragraph.add_run().add_break()
        run = paragraph.add_run(line)
        run.font.name = "Consolas"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.size = Pt(9.5)
    doc.add_paragraph()


def build() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("GEOFlow 会员系统方案")
    title_run.bold = True
    title_run.font.name = "Microsoft YaHei"
    title_run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    title_run.font.size = Pt(22)
    title_run.font.color.rgb = RGBColor(11, 37, 69)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run("多租户会员等级、用户分配、文章与知识库额度展示方案")
    subtitle_run.font.name = "Microsoft YaHei"
    subtitle_run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    subtitle_run.font.size = Pt(11)
    subtitle_run.font.color.rgb = RGBColor(85, 85, 85)

    add_heading(doc, "一、整体关系", 1)
    add_body(doc, "会员系统不替代现有角色权限系统，而是在现有多租户后台外增加一层会员套餐和额度控制。")
    add_bullets(doc, [
        "会员等级：套餐模板，例如入门版、专业版、商业版、企业版。",
        "用户/租户会员：某个用户或租户当前使用哪个会员等级，以及开始时间、到期时间。",
        "会员用量：当前用户或租户已经使用了多少文章额度，以及当前拥有多少知识库。",
        "角色仍然只表达权限，例如普通管理员和超级管理员；会员表达额度和有效期。",
    ])

    add_heading(doc, "二、超级管理员：会员等级管理", 1)
    add_body(doc, "入口放在后台左侧导航栏，新增“会员管理”。该页面只给超级管理员使用，用于维护会员等级和对应权益参数。第一版会员只限制每月发布文章数和知识库总数，不限制分发渠道和自动任务。")
    add_table(
        doc,
        ["等级名称", "每月文章数", "知识库总数", "状态", "操作"],
        [
            ["入门版", "50", "3", "启用", "编辑"],
            ["专业版", "300", "10", "启用", "编辑"],
            ["商业版", "1000", "30", "启用", "编辑"],
            ["企业版", "自定义", "自定义", "启用", "编辑"],
        ],
        [1.35, 1.45, 1.45, 1.05, 1.0],
    )
    add_body(doc, "新增或编辑会员等级时，弹窗字段包括等级名称、每月文章数、知识库总数和状态。")

    add_heading(doc, "三、超级管理员：用户管理里分配会员", 1)
    add_body(doc, "会员分配直接放在现有用户管理页面里实现。在“角色”展示位置下方增加一行会员信息，在操作区增加“分配会员”。")
    add_code_block(doc, [
        "账号信息                    角色/会员",
        "────────────────────────────────────",
        "zt970808zx@163.com          普通管理员",
        "zt970808zx                  专业版",
        "zt970808zx@163.com          到期：2026-08-03",
    ])
    add_body(doc, "不同状态下的展示文案：")
    add_table(
        doc,
        ["状态", "展示"],
        [
            ["未开通", "普通管理员 / 未开通会员"],
            ["正常", "普通管理员 / 专业版 / 到期：2026-08-03"],
            ["即将到期", "普通管理员 / 专业版 / 7 天后到期"],
            ["已到期", "普通管理员 / 专业版 / 已到期：2026-06-30"],
        ],
        [1.4, 5.1],
    )

    add_heading(doc, "四、分配会员弹窗", 1)
    add_body(doc, "点击用户管理里的“分配会员”后打开弹窗。为减少重复操作，弹窗提供固定周期选择，并自动计算到期时间。")
    add_code_block(doc, [
        "┌────────────────────────────────────┐",
        "│ 分配会员                           │",
        "├────────────────────────────────────┤",
        "│ 当前用户：zt970808zx                │",
        "│ 当前会员：专业版                    │",
        "│ 当前到期：2026-08-03                │",
        "│                                    │",
        "│ 会员等级：[ 专业版 ▼ ]              │",
        "│ 开通周期：[1个月] [1季度] [1年] [自定义] │",
        "│ 生效方式：从当前到期时间续期 / 从今天重新开通 │",
        "│ 开始时间：[ 2026-08-03 ]            │",
        "│ 到期时间：[ 2026-09-03 ]            │",
        "│ 备注：[ 手动续费 1 个月 ]            │",
        "│                                    │",
        "│ [取消]                    [保存]   │",
        "└────────────────────────────────────┘",
    ])
    add_table(
        doc,
        ["开通周期", "计算规则"],
        [
            ["1 个月", "开始时间 + 1 个月"],
            ["1 季度", "开始时间 + 3 个月"],
            ["1 年", "开始时间 + 12 个月"],
            ["自定义", "允许手动选择到期时间"],
        ],
        [1.4, 5.1],
    )
    add_bullets(doc, [
        "如果用户当前会员未过期，默认选择“从当前到期时间续期”。",
        "如果用户未开通或已到期，默认选择“从今天重新开通”。",
        "例如当前到期为 2026-08-03，选择 1 个月后，新到期时间为 2026-09-03。",
    ])

    add_heading(doc, "五、普通管理员：右上角会员摘要", 1)
    add_body(doc, "普通管理员只能查看自己的会员信息，不能修改会员等级。入口放在右上角个人下拉框。")
    add_code_block(doc, [
        "┌────────────────────────────────────┐",
        "│ 欢迎：zt970808zx                    │",
        "│ 普通管理员                          │",
        "├────────────────────────────────────┤",
        "│ 当前会员                 生效中     │",
        "│ 专业版                              │",
        "│ 到期时间：2026-08-03                │",
        "│ 本月文章：120 / 300                 │",
        "│ 知识库：6 / 10                      │",
        "│ [ 会员详情 ]                        │",
        "└────────────────────────────────────┘",
    ])

    add_heading(doc, "六、普通管理员：会员详情页", 1)
    add_body(doc, "从右上角下拉框点击“会员详情”进入，用于查看当前会员状态、到期时间和额度使用情况。额度只展示会员限制项：每月发布文章数和知识库总数。")
    add_table(
        doc,
        ["资源", "已用", "上限", "状态"],
        [
            ["本月发布文章", "120 篇", "300 篇", "正常"],
            ["知识库数量", "6 个", "10 个", "正常"],
        ],
        [2.1, 1.4, 1.4, 1.6],
    )
    add_body(doc, "未开通时展示：当前账号尚未开通会员，可查看已有数据，发布文章和创建知识库暂不可用，请联系管理员开通会员。")
    add_body(doc, "已到期时展示：会员已到期，部分新增功能已暂停，请联系管理员续费。")

    add_heading(doc, "七、后台顶部提醒", 1)
    add_body(doc, "提醒只在即将到期、已到期、未开通三种状态出现。文案保持温和，不使用红字，也不使用“禁止、错误、危险”等压迫性词。")
    add_table(
        doc,
        ["状态", "提醒文案"],
        [
            ["即将到期", "专业版会员将在 7 天后到期，请联系管理员续费。"],
            ["已到期", "会员已到期，发布文章和创建知识库已暂停。请联系管理员续费。"],
            ["未开通", "当前账号尚未开通会员，部分新增功能暂不可用。请联系管理员开通会员。"],
        ],
        [1.3, 5.2],
    )

    add_heading(doc, "八、第一版功能边界", 1)
    add_body(doc, "第一版重点做清楚会员管理、会员分配、会员展示和到期提醒，不引入支付和复杂商业化功能。")
    add_table(
        doc,
        ["第一版做", "第一版暂不做"],
        [
            ["会员等级管理", "在线支付"],
            ["用户管理里分配会员", "自动续费"],
            ["普通用户查看会员摘要", "优惠券和发票"],
            ["普通用户查看会员详情", "API 权益"],
            ["会员即将到期/已到期提醒", "分发渠道、自动任务、团队人数和复杂按量计费"],
        ],
        [3.25, 3.25],
    )

    add_heading(doc, "九、最终结论", 1)
    add_body(doc, "超级管理员通过左侧“会员管理”维护套餐参数，并在“用户管理”中给用户或租户分配会员。普通管理员通过右上角查看自己的会员状态和额度使用情况。会员即将到期或已到期时，在顶部和会员信息处做温和提醒。")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)


if __name__ == "__main__":
    build()
    print(OUTPUT)
